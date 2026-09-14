from datetime import date
from pathlib import Path

import docx

from medbot.extract import Block, Section
from medbot.docx_writer import write_docx

CFG = {"output": {"font": "Times New Roman", "font_size": 13}}


def sample_pick() -> dict:
    return {
        "source": "Healthline",
        "title": "Original Title",
        "title_vi": "Tiêu đề đã dịch",
        "url": "https://x.test/bai-goc",
        "published": "2026-09-13T10:00:00+00:00",
    }


def sample_sections() -> list[Section]:
    lead = Section(heading=None, blocks=[Block(kind="p", text="lead", text_vi="Đoạn mở đầu đã dịch")])
    body = Section(heading="H2 gốc", heading_vi="Mục con đã dịch", blocks=[
        Block(kind="p", text="p", text_vi="Đoạn văn thường."),
        Block(kind="list", items=["a", "b"], items_vi=["Mục một", "Mục hai"]),
    ])
    return [lead, body]


def test_title_becomes_heading_1(tmp_path: Path):
    out = tmp_path / "bai.docx"
    write_docx(out, sample_pick(), sample_sections(), CFG)
    doc = docx.Document(out)
    heading = next(p for p in doc.paragraphs if p.style.name == "Heading 1")
    assert heading.text == "Tiêu đề đã dịch"


def test_metadata_line_has_source_and_hyperlink(tmp_path: Path):
    out = tmp_path / "bai.docx"
    write_docx(out, sample_pick(), sample_sections(), CFG)
    doc = docx.Document(out)
    assert any("Healthline" in p.text for p in doc.paragraphs)
    rels = doc.part.rels
    assert any("bai-goc" in r.target_ref for r in rels.values() if r.reltype.endswith("hyperlink"))


def test_lead_paragraph_is_italic(tmp_path: Path):
    out = tmp_path / "bai.docx"
    write_docx(out, sample_pick(), sample_sections(), CFG)
    doc = docx.Document(out)
    sapo = next(p for p in doc.paragraphs if "Đoạn mở đầu đã dịch" in p.text)
    assert all(run.italic for run in sapo.runs if run.text.strip())


def test_h2_becomes_heading_2_with_translated_text(tmp_path: Path):
    out = tmp_path / "bai.docx"
    write_docx(out, sample_pick(), sample_sections(), CFG)
    doc = docx.Document(out)
    heading2 = next(p for p in doc.paragraphs if p.style.name == "Heading 2")
    assert heading2.text == "Mục con đã dịch"


def test_bullet_list_uses_translated_items(tmp_path: Path):
    out = tmp_path / "bai.docx"
    write_docx(out, sample_pick(), sample_sections(), CFG)
    doc = docx.Document(out)
    bullets = [p.text for p in doc.paragraphs if p.style.name == "List Bullet"]
    assert bullets == ["Mục một", "Mục hai"]


def test_footer_mentions_source_url(tmp_path: Path):
    out = tmp_path / "bai.docx"
    write_docx(out, sample_pick(), sample_sections(), CFG)
    doc = docx.Document(out)
    assert any("x.test/bai-goc" in p.text for p in doc.paragraphs)


def test_image_with_bytes_is_embedded(tmp_path: Path):
    png = bytes.fromhex(
        "89504e470d0a1a0a0000000d4948445200000001000000010802000000907753"
        "de0000000c4944415478da6360000002000155a2d4180000000049454e44ae426082"
    )
    section = Section(heading="Có ảnh", heading_vi="Có ảnh",
                       blocks=[Block(kind="image", src="https://x.test/a.jpg",
                                     alt_vi="Chú thích", image_bytes=png)])
    out = tmp_path / "bai.docx"
    write_docx(out, sample_pick(), [section], CFG)
    doc = docx.Document(out)
    assert len(doc.inline_shapes) == 1


def test_image_without_bytes_is_skipped_not_crashed(tmp_path: Path):
    section = Section(heading="Ảnh lỗi", heading_vi="Ảnh lỗi",
                       blocks=[Block(kind="image", src="https://x.test/a.jpg", image_bytes=None)])
    out = tmp_path / "bai.docx"
    write_docx(out, sample_pick(), [section], CFG)  # không được ném lỗi
    doc = docx.Document(out)
    assert len(doc.inline_shapes) == 0


def test_writes_to_onedrive_style_path_atomically(tmp_path: Path):
    nested = tmp_path / "OneDrive" / "2026-09"
    out = nested / "bai.docx"
    write_docx(out, sample_pick(), sample_sections(), CFG)
    assert out.exists()
    assert list(nested.glob("*.tmp*")) == []  # không để lại file tạm
