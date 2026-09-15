from __future__ import annotations

import io
import os
import tempfile
from datetime import date, datetime
from pathlib import Path

import docx
from docx.image.exceptions import (
    InvalidImageStreamError,
    UnexpectedEndOfFileError,
    UnrecognizedImageError,
)
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

_UNEMBEDDABLE_IMAGE_ERRORS = (
    UnrecognizedImageError,
    InvalidImageStreamError,
    UnexpectedEndOfFileError,
)

from medbot.extract import Block, Section


def add_hyperlink(paragraph, url: str, text: str) -> None:
    """python-docx không có hàm cấp cao cho hyperlink — phải dựng XML tay.

    Đã kiểm chứng thật: ghi ra rồi đọc lại bằng chính python-docx xác nhận
    quan hệ hyperlink tồn tại đúng trong file.
    """
    part = paragraph.part
    r_id = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    rpr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(underline)
    run.append(rpr)

    text_el = OxmlElement("w:t")
    text_el.text = text
    run.append(text_el)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _format_published(iso_string: str) -> str:
    return datetime.fromisoformat(iso_string).strftime("%d/%m")


def _write_block(doc, block: Block, italic: bool) -> None:
    if block.kind == "p":
        text = block.text_vi or block.text or ""
        if not text:
            return
        paragraph = doc.add_paragraph(text)
        if italic:
            for run in paragraph.runs:
                run.italic = True
    elif block.kind == "list":
        for item in (block.items_vi or block.items or []):
            doc.add_paragraph(item, style="List Bullet")
    elif block.kind == "image" and block.image_bytes:
        paragraph = doc.add_paragraph()
        try:
            paragraph.add_run().add_picture(io.BytesIO(block.image_bytes), width=Inches(5.5))
        except _UNEMBEDDABLE_IMAGE_ERRORS:
            # Ảnh đủ dung lượng qua bộ lọc rác của images.py nhưng
            # python-docx không nhận diện được định dạng — đã gặp thật với
            # JPEG "trần" (không có marker JFIF/Exif) từ CDN xử lý lại ảnh.
            # Bỏ qua ảnh này, không để hỏng cả bài (spec §10: ảnh lỗi thì
            # bỏ qua, không crash).
            paragraph._p.getparent().remove(paragraph._p)
            return
        caption_text = block.alt_vi or block.alt
        if caption_text:
            caption = doc.add_paragraph(caption_text)
            for run in caption.runs:
                run.italic = True


def write_docx(path: Path, pick: dict, sections: list[Section], cfg: dict) -> None:
    """Ghi bố cục §11: Heading 1 -> metadata (nguồn, ngày, hyperlink) ->
    sapo in nghiêng -> nội dung giữ cấp heading/bullet/ảnh -> chân trang.

    Ghi ra file tạm cùng thư mục rồi os.replace, cùng lý do với
    atomic_write_text ở kế hoạch 1: OneDrive có thể khoá file giữa chừng.
    """
    doc = docx.Document()
    style = doc.styles["Normal"]
    style.font.name = cfg["output"]["font"]
    style.font.size = Pt(cfg["output"]["font_size"])

    doc.add_heading(pick.get("title_vi") or pick["title"], level=1)

    meta = doc.add_paragraph()
    meta.add_run(f'{pick["source"]} · {_format_published(pick["published"])} · ')
    add_hyperlink(meta, pick["url"], "Xem bài gốc")

    for section in sections:
        if section.heading is None:
            for block in section.blocks:
                _write_block(doc, block, italic=True)
        else:
            doc.add_heading(section.heading_vi or section.heading, level=2)
            for block in section.blocks:
                _write_block(doc, block, italic=False)

    footer = doc.add_paragraph()
    footer_run = footer.add_run(
        f'Bản dịch tham khảo, thực hiện {date.today():%d/%m/%Y}. Nguồn: {pick["url"]}'
    )
    footer_run.italic = True

    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(dir=str(path.parent), suffix=".docx.tmp")
    os.close(fd)
    try:
        doc.save(tmp)
        os.replace(tmp, path)
    except BaseException:
        Path(tmp).unlink(missing_ok=True)
        raise
